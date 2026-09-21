"""Load workbooks into immutable cell records (formula + Excel's cached value)."""
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import xlrd
from openpyxl.utils import get_column_letter

XLS_EMPTY_TYPES = (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)


@dataclass(frozen=True)
class Cell:
    row: int
    col: int
    formula: str | None
    value: object
    unlocked: bool

    @property
    def coord(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"


@dataclass(frozen=True)
class Sheet:
    name: str
    cells: tuple[Cell, ...]
    validations: tuple[str, ...]
    protected: bool


def _formula_text(raw: object) -> str | None:
    text = getattr(raw, "text", raw)  # ArrayFormula keeps its body in .text
    if not isinstance(text, str) or not text.startswith("="):
        return None
    return f"{{{text}}}" if text is not raw else text


def _xls_value(cell: xlrd.sheet.Cell) -> object:
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, "#ERR")
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    return cell.value


def cached_values(path: Path) -> dict[str, dict[tuple[int, int], object]]:
    """Values as last computed by Excel, keyed by sheet then (row, col)."""
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(str(path))
        return {
            s.name: {
                (r + 1, c + 1): _xls_value(s.cell(r, c))
                for r in range(s.nrows)
                for c in range(s.ncols)
                if s.cell_type(r, c) not in XLS_EMPTY_TYPES
            }
            for s in book.sheets()
        }
    book = openpyxl.load_workbook(path, data_only=True)
    return {
        ws.title: {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
        for ws in book.worksheets
    }


def _validations(ws) -> tuple[str, ...]:
    return tuple(
        f"{dv.sqref} {dv.type}: {dv.formula1}"
        for dv in ws.data_validations.dataValidation
        if dv.formula1
    )


def load_sheets(formula_path: Path, original_path: Path) -> tuple[Sheet, ...]:
    """Formulas come from formula_path (xlsx); cached values from the original file."""
    values = cached_values(original_path)
    book = openpyxl.load_workbook(formula_path, data_only=False)
    sheets = []
    for ws in book.worksheets:
        sheet_values = values.get(ws.title, {})
        cells = tuple(
            Cell(
                row=c.row,
                col=c.column,
                formula=_formula_text(c.value),
                value=sheet_values.get((c.row, c.column)) if _formula_text(c.value) else c.value,
                unlocked=not c.protection.locked,
            )
            for row in ws.iter_rows()
            for c in row
            if c.value is not None
        )
        sheets.append(Sheet(ws.title, cells, _validations(ws), bool(ws.protection.sheet)))
    return tuple(sheets)
